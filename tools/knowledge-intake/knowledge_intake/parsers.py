from __future__ import annotations

import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from .common import Budget, ImportLimit, IntakeError, Limits, digest


class _HTML(HTMLParser):
    def __init__(self, limit: int):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.length = 0
        self.limit = limit
        self.hidden = 0
        self.pre = 0

    def add(self, text: str) -> None:
        self.length += len(text)
        if self.length > self.limit:
            raise ImportLimit("HTML output budget reached")
        self.parts.append(text)

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden += 1
        if self.hidden:
            return
        if re.fullmatch("h[1-6]", tag):
            self.add("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in {"p", "div", "section", "article", "tr"}:
            self.add("\n")
        elif tag == "br":
            self.add("\n")
        elif tag == "li":
            self.add("\n- ")
        elif tag in {"td", "th"}:
            self.add(" | ")
        elif tag == "pre":
            self.pre += 1
            self.add("\n\n````\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "template"} and self.hidden:
            self.hidden -= 1
            return
        if self.hidden:
            return
        if tag == "pre":
            self.pre = max(0, self.pre - 1)
            self.add("\n````\n")
        elif tag in {"p", "div", "section", "article", "tr"} or re.fullmatch("h[1-6]", tag):
            self.add("\n")

    def handle_data(self, text):
        if not self.hidden:
            self.add(text if self.pre else re.sub(r"\s+", " ", text))


def _archive_guard(path: Path, limits: Limits) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > limits.scan_entries or sum(entry.file_size for entry in entries) > limits.archive_bytes:
            raise ImportLimit("Office archive expansion budget reached")
        if any(entry.flag_bits & 1 for entry in entries):
            raise IntakeError("encrypted Office archives are unsupported")


def parse(path: Path, *, options: dict, limits: Limits, work_root: Path) -> dict:
    """Actual format parsing, executed in a bounded child process by sync."""
    budget = Budget(limits)
    parts, spans, assets, warnings = [], [], [], []
    count = 0
    next_line = 1
    partial = False

    def add(text: str, **location) -> None:
        nonlocal count, partial, next_line
        budget.check()
        text = text.strip()
        remaining = limits.output_chars - count
        if len(text) > remaining:
            text = text[:max(0, remaining)].rstrip()
            warnings.append("parsed output character budget reached")
            partial = True
        start = next_line
        parts.append(text)
        next_line += text.count("\n") + 2
        count += len(text) + 2
        spans.append({**location, "line_start": start, "line_end": start + text.count("\n"), "sha256": digest(text.encode("utf-8"))})
        if count >= limits.output_chars:
            raise ImportLimit("parsed output character budget reached")

    def image(image_path: Path, *, page: int | None = None) -> None:
        nonlocal partial
        from PIL import Image
        with Image.open(image_path) as opened:
            if opened.width * opened.height > limits.image_pixels:
                raise ImportLimit("image pixel budget reached")
            if getattr(opened, "n_frames", 1) > 1:
                warnings.append("only the first image frame is represented")
                partial = True
        asset = {"path": str(image_path.resolve()), "name": image_path.name, "sha256": digest(image_path.read_bytes())}
        if page:
            asset["page"] = page
        assets.append(asset)
        ocr = options.get("ocr")
        if ocr:
            try:
                from .ocr import recognize
                result = recognize(image_path, ocr if isinstance(ocr, dict) else {}, timeout=min(30, budget.remaining()), output_chars=limits.output_chars - count)
                if result["text"].strip():
                    add((f"## Page {page}: OCR\n\n" if page else "## OCR\n\n") + result["text"], page=page, kind="ocr", provider=result["provider"], asset_sha256=asset["sha256"])
                else:
                    warnings.append(f"OCR returned no text{f' on page {page}' if page else ''}")
            except (IntakeError, OSError, ValueError) as exc:
                warnings.append(f"OCR unavailable: {exc}")
                partial = True
        else:
            warnings.append(f"image text requires native caption or explicitly enabled OCR{f' (page {page})' if page else ''}")
            partial = True

    suffix = path.suffix.lower()
    try:
        if suffix in {".md", ".markdown", ".txt"}:
            add(path.read_text(encoding="utf-8-sig"), kind="text")
        elif suffix in {".html", ".htm"}:
            parser = _HTML(limits.output_chars)
            parser.feed(path.read_text(encoding="utf-8-sig"))
            add("".join(parser.parts), kind="html")
        elif suffix == ".docx":
            _archive_guard(path, limits)
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            for position, child in enumerate(document.element.body):
                budget.scan()
                if child.tag == qn("w:p"):
                    paragraph = Paragraph(child, document)
                    text = paragraph.text
                    name = paragraph.style.name if paragraph.style else ""
                    if name.startswith("Heading ") and name[-1:].isdigit():
                        text = "#" * min(int(name[-1]), 6) + " " + text
                    if text.strip():
                        add(text, block=position + 1, kind="paragraph")
                elif child.tag == qn("w:tbl"):
                    table = Table(child, document)
                    add("\n".join(" | ".join(cell.text.replace("\n", " ") for cell in row.cells) for row in table.rows), block=position + 1, kind="table")
            if document.inline_shapes:
                warnings.append("embedded Office images are retained in the original document; not OCR extracted")
                partial = True
        elif suffix == ".pptx":
            _archive_guard(path, limits)
            from pptx import Presentation
            presentation = Presentation(path)
            for number, slide in enumerate(presentation.slides, 1):
                if number > limits.pages:
                    raise ImportLimit("slide budget reached")
                text = [f"## Slide {number}"]
                for shape in slide.shapes:
                    budget.scan()
                    if shape.has_text_frame:
                        text.append(shape.text)
                    if shape.has_table:
                        text.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                    if shape.shape_type == 13:
                        warnings.append(f"slide {number} embedded image is retained in original presentation")
                        partial = True
                if slide.has_notes_slide:
                    text.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
                add("\n\n".join(text), slide=number, kind="slide")
        elif suffix == ".xlsx":
            _archive_guard(path, limits)
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            cells = 0
            try:
                for sheet in workbook:
                    budget.scan()
                    add(f"## Sheet: {sheet.title}", sheet=sheet.title, kind="heading")
                    for row in sheet.iter_rows():
                        budget.check()
                        cells += len(row)
                        if cells > limits.cells:
                            raise ImportLimit("spreadsheet cell budget reached")
                        values = [f"{cell.coordinate}: {cell.value}" for cell in row if cell.value is not None]
                        if values:
                            add(" | ".join(values), sheet=sheet.title, row=row[0].row, kind="cells")
            finally:
                workbook.close()
        elif suffix == ".pdf":
            import pypdfium2 as pdfium
            document = pdfium.PdfDocument(path)
            try:
                for index in range(min(len(document), limits.pages)):
                    budget.check()
                    page = document[index]
                    try:
                        textpage = page.get_textpage()
                        try:
                            if textpage.count_chars() > limits.output_chars - count:
                                raise ImportLimit("PDF text character budget reached")
                            text = textpage.get_text_range()
                        finally:
                            textpage.close()
                        if text.strip():
                            add(f"## Page {index + 1}\n\n{text}", page=index + 1, kind="pdf_text")
                        else:
                            width, height = page.get_size()
                            scale = min(2, (limits.image_pixels / max(1, width * height)) ** .5)
                            bitmap = page.render(scale=scale)
                            try:
                                page_path = work_root / f"page-{index + 1}.png"
                                pil = bitmap.to_pil()
                                try:
                                    pil.save(page_path)
                                finally:
                                    pil.close()
                            finally:
                                bitmap.close()
                            image(page_path, page=index + 1)
                    finally:
                        page.close()
                if len(document) > limits.pages:
                    raise ImportLimit("PDF page budget reached")
            finally:
                document.close()
        elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
            image(path)
        else:
            raise ValueError(f"unsupported format: {suffix}")
    except ImportLimit as exc:
        warnings.append(str(exc))
        partial = True
    except ImportError as exc:
        raise IntakeError(f"{suffix} parser is unavailable; install knowledge-intake[formats] ({exc.name})") from exc
    return {"text": "\n\n".join(parts), "spans": spans, "assets": assets, "warnings": list(dict.fromkeys(warnings)), "partial": partial, "parser": f"knowledge-intake/0.1:{suffix}"}
