"""Reproducible document/scan fixtures; values are synthetic, never benchmarks."""
from pathlib import Path
import argparse


def create(root: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas
    root.mkdir(parents=True, exist_ok=True)
    (root / "hccl.md").write_text("# HCCL communication reference\n\nSynthetic format fixture. Use current CANN/HCCL evidence before drawing conclusions.\n\n## Conditions\n\nAscend 910B; graph replay; no real benchmark run is represented.\n", encoding="utf-8")
    (root / "acl.html").write_text("<h1>ACL graph replay</h1><p>Synthetic parser fixture for vLLM-Ascend.</p><table><tr><th>Stage</th><th>Evidence</th></tr><tr><td>Capture</td><td>fixture only</td></tr></table>", encoding="utf-8")
    document = Document()
    document.add_heading("Ascend NPU report", 1)
    document.add_paragraph("Synthetic document: HCCL topology and ACL graph conditions.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Source", "CANN/HCCL documentation"
    document.save(root / "report.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "vLLM-Ascend graph replay"
    slide.placeholders[1].text = "Synthetic slide\nNo performance measurement"
    presentation.save(root / "graph.pptx")
    workbook = Workbook()
    workbook.active.title = "Synthetic"
    workbook.active.append(["NPU", "fixture throughput (not measured)"])
    workbook.active.append(["910B", 120])
    workbook.active.append(["910C", 180])
    workbook.save(root / "data.xlsx")
    pdf = canvas.Canvas(str(root / "text.pdf"))
    pdf.drawString(72, 700, "HCCL / ACL graph reference: synthetic text PDF fixture")
    pdf.save()
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
        small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 24)
    except OSError:
        font = small = ImageFont.load_default(size=26)
    draw.text((45, 25), "Synthetic HCCL graph replay chart", fill="black", font=font)
    draw.text((45, 75), "Illustration only - not measured performance", fill="black", font=small)
    draw.line((150, 160, 150, 560), fill="black", width=3)
    draw.line((150, 560, 1080, 560), fill="black", width=3)
    draw.text((30, 130), "tokens/s", fill="black", font=small)
    for value in (0, 60, 120, 180):
        y = 560 - value * 2
        draw.text((75, y - 12), str(value), fill="black", font=small)
        draw.line((145, y, 1080, y), fill="#cccccc", width=1)
    draw.rectangle((300, 320, 500, 560), fill="#3973b9")
    draw.rectangle((690, 200, 890, 560), fill="#269177")
    draw.text((350, 275), "120", fill="black", font=font)
    draw.text((750, 155), "180", fill="black", font=font)
    draw.text((310, 580), "910B fixture", fill="black", font=small)
    draw.text((700, 580), "910C fixture", fill="black", font=small)
    draw.text((45, 645), "Batch and software version intentionally unspecified", fill="black", font=small)
    image.save(root / "graph-chart.png")
    scan = canvas.Canvas(str(root / "scan.pdf"), pagesize=(600, 350))
    scan.drawImage(str(root / "graph-chart.png"), 0, 0, width=600, height=350)
    scan.save()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    create(parser.parse_args().directory)
